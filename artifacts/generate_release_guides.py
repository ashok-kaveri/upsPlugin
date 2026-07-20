from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    ListFlowable,
    ListItem,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ARTIFACTS_DIR = Path(__file__).resolve().parent
GENERATED_ON = "June 16, 2026"


@dataclass
class Guide:
    lane: str
    version: str
    product: str
    summary: str
    cards: list[str]
    changes: list[str]
    customer_impact: list[str]
    support_points: list[str]
    verification: list[str]
    scenarios: list[str]
    escalation: list[str]
    references: list[str]
    basename: str


GUIDES = [
    Guide(
        lane="FedEx Release - v8.3.2",
        version="v8.3.2",
        product="FedEx Shipping Plugin for WooCommerce",
        summary="This release focuses on label accuracy, time-sensitive pickup handling, and clearer separation between FedEx registration details and general plugin settings.",
        cards=[
            "Void Shipment + Swapped label address",
            "Fix Express pickup date to advance when ready time has passed",
            "Separate registration details from settings options",
        ],
        changes=[
            "Shipment voiding and label-address handling were updated so support can better trust the shipment state and printed address details after corrections or cancellations.",
            "Express pickup-date selection now advances when the configured ready time has already passed, helping avoid same-day pickup requests that are no longer valid.",
            "Registration details were separated from broader settings so merchants can more clearly manage onboarding data without confusing it with day-to-day shipping options.",
        ],
        customer_impact=[
            "Merchants should see fewer cases where label details appear inconsistent after shipment voiding or rework.",
            "Pickup scheduling should behave more predictably when labels are created later in the day.",
            "Carrier setup screens should feel clearer, especially for merchants revisiting credentials or registration-related values.",
        ],
        support_points=[
            "If a merchant reports swapped or unexpected label address details, confirm whether the issue happened before or after a void/retry workflow.",
            "If a same-day Express pickup is not offered after the ready time has passed, explain that the plugin now shifts to the next valid pickup date instead of sending an outdated request.",
            "If a merchant cannot find a registration field, guide them to the registration-specific section rather than the general shipping settings area.",
        ],
        verification=[
            "Create and void a test shipment, then verify the follow-up label reflects the correct shipment address details.",
            "Set an Express-ready time earlier than the current time and confirm the pickup date advances to the next valid option.",
            "Open settings and confirm registration-related values are shown separately from standard shipping preferences.",
        ],
        scenarios=[
            "A merchant voids a shipment and recreates the label after correcting an address.",
            "A merchant creates an Express shipment after the daily ready-time cutoff.",
            "A merchant revisits FedEx setup and needs to update registration details without changing regular shipping behavior.",
        ],
        escalation=[
            "Escalate if label addresses are still swapped after reproducing the issue with a fresh shipment.",
            "Escalate if the pickup date does not advance and FedEx rejects the request due to time cutoff issues.",
            "Escalate UI placement issues if required registration fields are missing or inaccessible.",
        ],
        references=[
            "Release lane: FedEx Release - v8.3.2",
            "Source cards: Void Shipment + Swapped label address; Fix Express pickup date to advance when ready time has passed; Separate registration details from settings options",
        ],
        basename="FedEx_Release_v8.3.2_Support_Guide",
    ),
    Guide(
        lane="FedEx Release - v8.4.0",
        version="v8.4.0",
        product="FedEx Shipping Plugin for WooCommerce",
        summary="This release expands customs and invoice controls for international shipping, especially around tax identifiers, importer details, commercial invoice fields, and currency handling.",
        cards=[
            "Added an option to use the order currency for customs and declared value on individual orders, with the setting now available under Advanced settings.",
            "Consignee / Billing Tax ID at Checkout — collect customer VAT/EORI/TIN",
            "Importer of Record for International Shipments",
            "Order-Level Commercial Invoice Fields — Purpose, Comments & Special Instructions",
            "Currency Conversion Rounding Causing Declared Value to Exceed Customs Value",
        ],
        changes=[
            "A new advanced option allows order currency to be used for customs and declared-value handling on individual orders.",
            "Checkout can now collect consignee or billing tax identifiers such as VAT, EORI, or TIN when merchants need them for compliance workflows.",
            "Importer-of-record handling was added for international shipments that require a distinct importing party.",
            "Order-level commercial invoice fields now support purpose, comments, and special instructions.",
            "Currency-conversion rounding was corrected so declared value does not exceed customs value because of conversion math.",
        ],
        customer_impact=[
            "International merchants should have more complete customs data available before label generation.",
            "Stores shipping across borders may see fewer manual edits to invoice and customs details.",
            "Merchants using different order currencies should see more consistent declared/customs values.",
        ],
        support_points=[
            "If a merchant wants customs amounts to follow order currency, direct them to the new advanced setting and confirm the order-level expectation.",
            "If a merchant asks where to collect VAT, EORI, or TIN, explain that checkout support now exists for consignee or billing tax IDs.",
            "If customs paperwork requires an importer different from shipper or recipient, verify whether Importer of Record is now configured.",
            "If merchants ask for more invoice context, point them to the order-level commercial invoice fields for purpose, comments, and special instructions.",
            "If declared value previously exceeded customs value by a small amount, note that rounding behavior was corrected in this release.",
        ],
        verification=[
            "Place an international test order in a non-default order currency and confirm customs/declared values follow the configured behavior.",
            "Confirm tax-ID fields can be collected and passed when a checkout flow requires VAT, EORI, or TIN.",
            "Create an international shipment that uses Importer of Record and verify the related details appear in the shipping workflow.",
            "Populate purpose, comments, and special instructions on an order and verify they appear in the commercial invoice flow.",
            "Check a currency-conversion case that previously over-rounded declared value and confirm the values now stay aligned.",
        ],
        scenarios=[
            "A merchant ships internationally and needs customs values in the same currency as the order.",
            "A store must collect customer tax IDs for regulatory or clearance requirements.",
            "An importer is legally different from the shipper or receiver on the order.",
            "The merchant needs shipment-specific invoice instructions beyond default product data.",
        ],
        escalation=[
            "Escalate if customs values still exceed expected limits after confirming currency settings.",
            "Escalate if tax-ID data is collected at checkout but does not reach the shipment workflow.",
            "Escalate if Importer of Record or invoice fields fail to appear on the generated documentation.",
        ],
        references=[
            "Release lane: FedEx Release - v8.4.0",
            "Source cards: order currency for customs and declared value; Consignee / Billing Tax ID; Importer of Record; Order-Level Commercial Invoice Fields; Currency Conversion Rounding",
        ],
        basename="FedEx_Release_v8.4.0_Support_Guide",
    ),
    Guide(
        lane="UPS Release - 6.6.0",
        version="6.6.0",
        product="UPS Shipping Plugin for WooCommerce",
        summary="This release covers invoice suppression for eligible EU shipments, checkout and label handling for signature and recipient TIN flows, Access Point block-checkout support, UI validation for per-package workflows, and additional handling support.",
        cards=[
            "Skip Commercial Invoice for EU Shipments",
            "Signature requirement breaks domestic checkout rates and is missing from US↔Puerto Rico labels",
            "Access Point Locator: Block Checkout Support",
            "Recipient TIN: Block Checkout Support + Configurable Label",
            "Per-Package UI Redesign: Complete UI/UX Testing",
            "Additional Handling Indicator",
        ],
        changes=[
            "Commercial invoices are skipped for eligible EU shipments where they should not be produced.",
            "Signature-requirement handling was updated for domestic checkout rates and US to Puerto Rico label scenarios.",
            "UPS Access Point Locator support was extended to WooCommerce block checkout.",
            "Recipient TIN support now covers block checkout and configurable label handling.",
            "The per-package UI redesign was validated for workflow and usability coverage.",
            "Additional Handling indicator support was added where relevant.",
        ],
        customer_impact=[
            "EU merchants should avoid unnecessary commercial invoices on shipments that do not require them.",
            "UPS signature options should behave more reliably during rating and label generation for the affected routes.",
            "Stores using block checkout gain better compatibility for Access Point and recipient tax workflows.",
            "Package-level workflows should feel clearer after the UI redesign validation.",
        ],
        support_points=[
            "If a merchant expects a commercial invoice for an EU shipment, first confirm whether that route should now intentionally skip it.",
            "If domestic UPS rates disappear after enabling signature requirements, verify whether the order matches the previously affected scenario.",
            "If an Access Point option is missing in block checkout, confirm the store is actually using the blocks-based checkout flow.",
            "If a merchant needs a recipient TIN on the label, check both the block-checkout data path and the label configuration.",
            "If the package UI looks different, explain that the per-package experience was redesigned and validated in this release.",
            "If shipment requirements mention special handling, verify whether the new Additional Handling indicator is expected in that flow.",
        ],
        verification=[
            "Create an eligible EU shipment and confirm no commercial invoice is generated when it should be skipped.",
            "Test a domestic UPS checkout flow with signature requirement enabled and confirm rates still appear correctly.",
            "Test a US to Puerto Rico label flow and verify the expected signature behavior is present on the label.",
            "Use WooCommerce block checkout to verify Access Point selection appears and behaves correctly.",
            "Verify recipient TIN entry and label output in a block-checkout scenario.",
            "Walk through package-level configuration to confirm the redesigned UI remains clear and functional.",
            "Confirm Additional Handling is shown or transmitted in a shipment that requires it.",
        ],
        scenarios=[
            "An EU merchant reports a missing invoice that is now intentionally suppressed.",
            "A UPS domestic shipment loses rates when signature requirement is enabled.",
            "A block-checkout store needs Access Point selection or recipient TIN support.",
            "A merchant needs extra confidence around package-level configuration after the UI update.",
        ],
        escalation=[
            "Escalate if EU shipments still generate unnecessary invoices after confirming route eligibility.",
            "Escalate if signature settings still remove rates or fail to appear on affected labels.",
            "Escalate if Access Point or recipient TIN support does not appear in block checkout despite valid setup.",
            "Escalate if Additional Handling is required but not reflected in the request or label output.",
        ],
        references=[
            "Release lane: UPS Release - 6.6.0",
            "Source cards: Skip Commercial Invoice for EU Shipments; Signature requirement breaks domestic checkout rates and is missing from US↔Puerto Rico labels; Access Point Locator; Recipient TIN; Per-Package UI Redesign; Additional Handling Indicator",
        ],
        basename="UPS_Release_v6.6.0_Support_Guide",
    ),
    Guide(
        lane="Shipment Tracking Release 3.3.0",
        version="3.3.0",
        product="Shipment Tracking Plugin for WooCommerce",
        summary="This release adds Canpar courier shipment-tracking integration to the Shipment Tracking plugin.",
        cards=[
            "Canpar courier - shipment tracking integration",
        ],
        changes=[
            "Canpar courier support was added to the shipment-tracking integration set.",
        ],
        customer_impact=[
            "Merchants using Canpar can now manage tracking through the Shipment Tracking plugin instead of relying on custom workarounds.",
            "Customers should receive a more consistent tracking experience when Canpar is the active carrier.",
        ],
        support_points=[
            "If a merchant asks whether Canpar is supported, confirm that native shipment-tracking integration is included in version 3.3.0.",
            "If tracking is not updating, first confirm the merchant selected Canpar in the tracking workflow and is using the updated plugin version.",
        ],
        verification=[
            "Add a shipment with Canpar as the carrier and verify tracking details can be saved and displayed correctly.",
            "Confirm the customer-facing tracking flow behaves as expected for a Canpar order.",
        ],
        scenarios=[
            "A merchant wants to enable Canpar without using a custom carrier workaround.",
            "Support needs to validate whether Canpar tracking details are being stored and shown properly.",
        ],
        escalation=[
            "Escalate if Canpar cannot be selected, saved, or rendered correctly after confirming version 3.3.0 is active.",
            "Escalate if customer-facing tracking links or status details fail specifically for Canpar while other carriers work.",
        ],
        references=[
            "Release lane: Shipment Tracking Release 3.3.0",
            "Source card: Canpar courier - shipment tracking integration",
        ],
        basename="Shipment_Tracking_Release_v3.3.0_Support_Guide",
    ),
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def ensure_styles(doc: Document) -> None:
    styles = doc.styles
    if "Guide Title" not in styles:
        title_style = styles.add_style("Guide Title", WD_STYLE_TYPE.PARAGRAPH)
        title_style.font.name = "Arial"
        title_style.font.size = Pt(22)
        title_style.font.bold = True
        title_style.font.color.rgb = RGBColor(24, 45, 84)
    if "Guide Heading" not in styles:
        heading_style = styles.add_style("Guide Heading", WD_STYLE_TYPE.PARAGRAPH)
        heading_style.font.name = "Arial"
        heading_style.font.size = Pt(13)
        heading_style.font.bold = True
        heading_style.font.color.rgb = RGBColor(24, 45, 84)
    if "Guide Body" not in styles:
        body_style = styles.add_style("Guide Body", WD_STYLE_TYPE.PARAGRAPH)
        body_style.font.name = "Arial"
        body_style.font.size = Pt(10.5)


def add_bullets_docx(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="Guide Body")
        p.style = doc.styles["List Bullet"]
        run = p.add_run(item)
        run.font.name = "Arial"
        run.font.size = Pt(10.5)


def build_docx(guide: Guide, path: Path) -> None:
    doc = Document()
    ensure_styles(doc)
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    title = doc.add_paragraph(style="Guide Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(f"{guide.product} Support Guide")

    subtitle = doc.add_paragraph(style="Guide Body")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run(f"{guide.lane} | Generated {GENERATED_ON}")
    subtitle_run.italic = True

    doc.add_paragraph("")

    table = doc.add_table(rows=3, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    metadata = [
        ("Release Lane", guide.lane),
        ("Plugin Version", guide.version),
        ("Primary Audience", "Customer support, QA, and implementation teams"),
    ]
    for row, (label, value) in zip(table.rows, metadata):
        row.cells[0].text = label
        row.cells[1].text = value
        set_cell_shading(row.cells[0], "DCE6F2")
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(10)
            row.cells[0].paragraphs[0].runs[0].bold = True

    sections = [
        ("Release Summary", [guide.summary]),
        ("Included Release Items", guide.cards),
        ("What Changed", guide.changes),
        ("Customer Impact", guide.customer_impact),
        ("Support Talking Points", guide.support_points),
        ("How Support Can Verify The Release", guide.verification),
        ("Expected Scenarios", guide.scenarios),
        ("Escalation Guidance", guide.escalation),
        ("Internal References", guide.references),
    ]

    for heading, items in sections:
        doc.add_paragraph("")
        p = doc.add_paragraph(style="Guide Heading")
        p.add_run(heading)
        add_bullets_docx(doc, items)

    doc.save(path)


def bullet_list_pdf(items: list[str], style: ParagraphStyle) -> ListFlowable:
    return ListFlowable(
        [ListItem(Paragraph(item, style)) for item in items],
        bulletType="bullet",
        leftIndent=18,
        bulletFontName="Helvetica",
        bulletFontSize=9,
    )


def build_pdf(guide: Guide, path: Path) -> None:
    doc = SimpleDocTemplate(
        str(path),
        pagesize=letter,
        leftMargin=0.65 * inch,
        rightMargin=0.65 * inch,
        topMargin=0.6 * inch,
        bottomMargin=0.6 * inch,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "GuideTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=24,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#182d54"),
        spaceAfter=6,
    )
    body_style = ParagraphStyle(
        "GuideBody",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=9.5,
        leading=13,
        spaceAfter=4,
    )
    heading_style = ParagraphStyle(
        "GuideHeading",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=14,
        textColor=colors.HexColor("#182d54"),
        spaceBefore=8,
        spaceAfter=6,
    )
    subtitle_style = ParagraphStyle(
        "GuideSubtitle",
        parent=body_style,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#555555"),
        spaceAfter=12,
    )

    story = [
        Paragraph(f"{guide.product} Support Guide", title_style),
        Paragraph(f"{guide.lane} | Generated {GENERATED_ON}", subtitle_style),
    ]

    metadata_table = Table(
        [
            ["Release Lane", guide.lane],
            ["Plugin Version", guide.version],
            ["Primary Audience", "Customer support, QA, and implementation teams"],
        ],
        colWidths=[1.55 * inch, 4.95 * inch],
    )
    metadata_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#DCE6F2")),
                ("GRID", (0, 0), (-1, -1), 0.6, colors.HexColor("#AAB7C4")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                ("FONTNAME", (1, 0), (1, -1), "Helvetica"),
                ("FONTSIZE", (0, 0), (-1, -1), 9.5),
                ("LEADING", (0, 0), (-1, -1), 12),
                ("TOPPADDING", (0, 0), (-1, -1), 7),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ]
        )
    )
    story.extend([metadata_table, Spacer(1, 0.14 * inch)])

    sections = [
        ("Release Summary", [guide.summary]),
        ("Included Release Items", guide.cards),
        ("What Changed", guide.changes),
        ("Customer Impact", guide.customer_impact),
        ("Support Talking Points", guide.support_points),
        ("How Support Can Verify The Release", guide.verification),
        ("Expected Scenarios", guide.scenarios),
        ("Escalation Guidance", guide.escalation),
        ("Internal References", guide.references),
    ]

    for heading, items in sections:
        story.append(Paragraph(heading, heading_style))
        story.append(bullet_list_pdf(items, body_style))

    doc.build(story)


def main() -> None:
    for guide in GUIDES:
        docx_path = ARTIFACTS_DIR / f"{guide.basename}.docx"
        pdf_path = ARTIFACTS_DIR / f"{guide.basename}.pdf"
        build_docx(guide, docx_path)
        build_pdf(guide, pdf_path)
        print(f"Generated {docx_path.name} and {pdf_path.name}")


if __name__ == "__main__":
    main()
